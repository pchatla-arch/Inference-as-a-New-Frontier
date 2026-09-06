#!/usr/bin/env python3
"""Build a native-equation editorial DOCX and readable Markdown from canonical TeX.

Requires a completed PDF build (paper/main.aux), Pandoc, and python-docx.
The IEEE PDF remains the layout reference; this export is intentionally single-column.
"""
from __future__ import annotations
import json
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'paper'


def pandoc(text: str, source: str='latex', extra: list[str] | None=None) -> dict[str, Any]:
    p = subprocess.run(['pandoc','-f',source,'-t','json',*(extra or [])], input=text,
                       text=True, cwd=PAPER, capture_output=True, check=True)
    return json.loads(p.stdout)


def braced(text: str, command: str) -> str:
    m = re.search(re.escape(command)+r'\s*\{',text)
    if not m:
        raise ValueError(f'Missing {command}')
    start=m.end(); depth=1
    for pos in range(start,len(text)):
        if text[pos]=='{' and text[pos-1]!='\\': depth+=1
        elif text[pos]=='}' and text[pos-1]!='\\':
            depth-=1
            if not depth:return text[start:pos]
    raise ValueError(f'Unclosed {command}')


def roman(n: int) -> str:
    out=''
    for v,t in [(1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')]:
        while n>=v:out+=t;n-=v
    return out


def main() -> None:
    tex=(PAPER/'main.tex').read_text()
    aux=(PAPER/'main.aux').read_text()
    refs=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',aux))
    keys=re.findall(r'\\abx@aux@defaultrefcontext\{0\}\{([^}]+)\}',aux)
    if not keys:raise RuntimeError('Run make paper before make editorial.')
    citations={k:i+1 for i,k in enumerate(keys)}
    title=braced(tex,'\\title')
    body=tex.split('\\begin{document}',1)[1].split('\\end{document}',1)[0]
    body=body.replace('\\maketitle','').replace('\\balance','').replace('\\printbibliography','')
    body=re.sub(r'\\IEEEPARstart\{([^}]+)\}\{([^}]+)\}',lambda m:m[1]+m[2],body)
    body=body.replace('\\begin{abstract}','\\section*{Abstract}').replace('\\end{abstract}','')
    body=body.replace('\\begin{IEEEkeywords}','\\textbf{Index Terms---}').replace('\\end{IEEEkeywords}','')
    body=re.sub(r'\\cite\{([^}]+)\}',lambda m:', '.join(f'[{citations[k.strip()]}]' for k in m[1].split(',')),body)
    body=re.sub(r'\\ref\{([^}]+)\}',lambda m:refs[m[1]],body)
    body=body.replace('^*', r'^{\text{*}}')
    body=re.sub(r'\\SI\{([^}]+)\}\{([^}]+)\}',r'\1~\2',body)
    body=re.sub(r'\\si\{([^}]+)\}',r'\1',body)
    fcount=0
    def figure(m):
        nonlocal fcount
        text=m[0];fcount+=1
        name=re.search(r'\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}',text)[1]
        name=Path(name).stem+'.png'
        caption=braced(text,'\\caption')
        width='6.3in' if '\\begin{figure*}' in text else '4.4in'
        return '\n\\begin{center}\n\\includegraphics[width='+width+']{figures_png/'+name+'}\n\\end{center}\n\\textit{Figure '+str(fcount)+'. '+caption+'}\n'
    body=re.sub(r'\\begin\{figure\*?\}.*?\\end\{figure\*?\}',figure,body,flags=re.S)
    tcount=0
    def table(m):
        nonlocal tcount
        text=m[0];tcount+=1
        caption=braced(text,'\\caption')
        data=text.split('\\toprule',1)[1].split('\\bottomrule',1)[0]
        data=data.replace('\\midrule','').replace('\\hline','')
        header=data.split('\\\\',1)[0]
        cols=header.count('&')+1
        return '\n\\textbf{Table '+roman(tcount)+'. '+caption+'}\n\\begin{longtable}{'+('l'*cols)+'}\n'+data+'\n\\end{longtable}\n'
    body=re.sub(r'\\begin\{table\*?\}.*?\\end\{table\*?\}',table,body,flags=re.S)
    eq=0
    def equation(m):
        nonlocal eq
        typ=m[1]; content=m[2]
        content=re.sub(r'\\label\{[^}]+\}','',content).strip()
        if typ=='align':
            parts=re.split(r'\\\\',content)
            out=[]
            for part in parts:
                if not part.strip():continue
                eq+=1;out.append('\\[\n'+part.strip().replace('&','')+'\n\\]\n\\texttt{EQLABEL'+str(eq)+'}\n')
            return '\n'.join(out)
        eq+=1
        if r'\begin{aligned}' in content:
            content=content.replace(r'\begin{aligned}','').replace(r'\end{aligned}','')
            rows=[part.strip().replace('&','') for part in content.split(chr(92)*2) if part.strip()]
            if len(rows)<=2:
                # Full-width Word does not need the IEEE column's line break.
                content=' '.join(rows)
            else:
                # Independent OfficeMath paragraphs avoid LibreOffice's eqArr import defect.
                chunks=[]
                for i,row in enumerate(rows):
                    chunks.append('\\[\n'+row+'\n\\]\n')
                    if i==0:chunks.append('\\texttt{EQLABEL'+str(eq)+'}\n')
                return '\n'.join(chunks)
        return '\\[\n'+content+'\n\\]\n\\texttt{EQLABEL'+str(eq)+'}\n'

    body=re.sub(r'\\begin\{(equation|align)\}(.*?)\\end\{\1\}',equation,body,flags=re.S)
    body=re.sub(r'\\label\{[^}]+\}','',body)
    preamble=tex.split('\\begin{document}')[0]
    norm=preamble+'\n\\begin{document}\n'+body+'\n\\end{document}'
    doc=pandoc(norm)
    # Add numbered headings directly: numbering remains identical across PDF and Word.
    sec=0; sub=0
    for b in doc['blocks']:
        if b['t']=='Header':
            level,attr,ins=b['c']
            txt=''.join(x.get('c',' ') if x['t']=='Str' else ' ' for x in ins)
            if txt in ('Abstract','Artifact Availability'):continue
            if level==1:
                sec+=1;sub=0; prefix=roman(sec)+'. '
            else:
                sub+=1;prefix=chr(64+sub)+'. '
            b['c'][2]=[{'t':'Str','c':prefix}]+ins
    # Citeproc only produces the reference list; canonical numeric labels above come from Biber.
    bibdoc=pandoc('\n\n'.join('[@'+k+']' for k in keys), 'markdown',
       ['--citeproc','--bibliography=references.bib','--csl=../scripts/ieee.csl'])
    refblocks=[]
    for b in bibdoc['blocks']:
        if b['t']=='Div' and b['c'][0][0]=='refs':refblocks=b['c'][1]
    if not refblocks:raise RuntimeError('Citeproc returned no bibliography')
    doc['blocks'] += [{'t':'Header','c':[1,['references',[],[]],[{'t':'Str','c':'References'}]]}]+refblocks
    doc['meta']={'title':{'t':'MetaInlines','c':[{'t':'Str','c':title}]},
                 'author':{'t':'MetaList','c':[{'t':'MetaInlines','c':[{'t':'Str','c':'Prasuna Chatla'}]}]},
                 'date':{'t':'MetaInlines','c':[{'t':'Str','c':'Independent Researcher | September 2026'}]}}
    with tempfile.TemporaryDirectory() as td:
        jp=Path(td)/'editorial.json';jp.write_text(json.dumps(doc))
        out=PAPER/'inference_as_a_new_frontier.docx'
        subprocess.run(['pandoc',str(jp),'-f','json','-o',str(out),'--resource-path='+str(PAPER)],check=True,cwd=PAPER)
        # Markdown copy has GitHub-renderable mathematics and relative image paths.
        md=json.loads(json.dumps(doc))
        # Convert equation marker to a visible label, not an internal export marker.
        def clean(x):
            if isinstance(x,dict):
                if x.get('t')=='Code' and re.fullmatch('EQLABEL[0-9]+',x.get('c',[None,''])[1]):
                    return {'t':'Str','c':'Equation ('+x['c'][1][7:]+')'}
                return {k:clean(v) for k,v in x.items()}
            if isinstance(x,list):return [clean(v) for v in x]
            return x
        mp=Path(td)/'markdown.json';mp.write_text(json.dumps(clean(md)))
        subprocess.run(['pandoc',str(mp),'-f','json','-t','gfm+tex_math_dollars','--wrap=none','-s','-o',str(PAPER/'paper.md')],check=True,cwd=PAPER)
    md_path=PAPER/'paper.md'
    md_text=re.sub(r'^---\n.*?\n---\n', '', md_path.read_text(), count=1, flags=re.S)
    md_path.write_text('# '+title+'\n\n**Prasuna Chatla** · Independent Researcher · September 2026\n\n'+md_text.lstrip())
    style_docx(out,title)
    print(f'Wrote {out.name}; {eq} equations, {fcount} figures, {tcount} tables, {len(keys)} references.')


def style_docx(path: Path, title: str) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document(path)
    for sec in d.sections:
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7)
        sec.left_margin=sec.right_margin=Inches(.85)
        sec.header_distance=sec.footer_distance=Inches(.3)
    styles={s.name:s for s in d.styles}
    for name in ['Normal','Body Text','First Paragraph','Compact']:
        if name in styles:
            st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(10.5)
            st.paragraph_format.space_after=Pt(6)
            st.paragraph_format.line_spacing=1.06
            st.paragraph_format.widow_control=True
    for name,size in [('Title',20),('Subtitle',11),('Heading 1',13),('Heading 2',11)]:
        st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('182F45');st.font.bold=name not in ['Subtitle']
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(12 if name=='Heading 1' else 8)
        st.paragraph_format.space_after=Pt(5)
    styles['Title'].paragraph_format.space_after=Pt(9)
    styles['Caption'].font.name='Times New Roman';styles['Caption'].font.size=Pt(9)
    styles['Caption'].paragraph_format.space_after=Pt(7)
    # Equations retain native OfficeMath, with a consistent right-side equation number.
    for p in list(d.paragraphs):
        if re.fullmatch(r'EQLABEL\d+',p.text):
            num=p.text[7:]
            prev=p._p.getprevious()
            if prev is None:raise RuntimeError('Equation label without equation')
            from docx.text.paragraph import Paragraph
            ep=Paragraph(prev,p._parent)
            ep.paragraph_format.keep_together=True
            ep.paragraph_format.space_before=Pt(5);ep.paragraph_format.space_after=Pt(7)
            mathparas=ep._p.findall(qn('m:oMathPara'))
            for mpara in mathparas:
                for math in list(mpara.findall(qn('m:oMath'))):
                    mpara.remove(math);ep._p.insert(list(ep._p).index(mpara),math)
                ep._p.remove(mpara)
            ep.paragraph_format.tab_stops.add_tab_stop(Inches(3.05),WD_TAB_ALIGNMENT.CENTER)
            ep.paragraph_format.tab_stops.add_tab_stop(Inches(6.75),WD_TAB_ALIGNMENT.RIGHT)
            run=OxmlElement('w:r');tab=OxmlElement('w:tab');run.append(tab)
            ep._p.insert(1 if ep._p.find(qn('w:pPr')) is not None else 0,run)
            ep.add_run('\t('+num+')')
            ep.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p._p.getparent().remove(p._p)
    for p in d.paragraphs:
        if p.text.startswith('Figure '):
            p.style=styles['Caption'];p.paragraph_format.keep_with_next=False
        if p.text.startswith('Table '):
            p.style=styles['Caption'];p.paragraph_format.keep_with_next=True
            for r in p.runs:r.bold=True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name in ['Title','Author','Date']:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name=='Bibliography':
            p.paragraph_format.space_after=Pt(4)
            for r in p.runs:r.font.size=Pt(9)
    for table in d.tables:
        table.autofit=False
        n=len(table.columns)
        widths={3:[1.32,2.65,2.73],4:[1.15,2.05,1.7,1.8]}.get(n,[6.7/n]*n)
        for col,w in zip(table.columns,widths):col.width=Inches(w)
        for ri,row in enumerate(table.rows):
            for ci,c in enumerate(row.cells):
                c.width=Inches(widths[ci])
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.line_spacing=1
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
                    if ri==0:
                        for r in p.runs:r.bold=True
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF4');c._tc.get_or_add_tcPr().append(shade)
            trpr=row._tr.get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if ri==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for s in d.sections:
        h=s.header.paragraphs[0];h.text='INFERENCE AS A NEW FRONTIER  |  PRASUNA CHATLA'
        h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for r in h.runs:r.font.size=Pt(8);r.font.name='Times New Roman'
        f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f.add_run('Technical preprint  |  ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
        for r in f.runs:r.font.size=Pt(8)
    for p in d.paragraphs:
        if p.text=='References':p.paragraph_format.page_break_before=True
    d.core_properties.title=title;d.core_properties.author='Prasuna Chatla'
    d.core_properties.subject='Data movement, LLM serving, RL rollouts, and evaluator-aware self-improvement'
    d.core_properties.comments='Editable editorial companion to the canonical IEEE-style PDF.'
    d.save(path)

if __name__=='__main__':main()
